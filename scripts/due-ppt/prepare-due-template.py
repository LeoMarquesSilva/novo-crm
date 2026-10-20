"""
Prepara o modelo de apresentação da Due Diligence para o gerador do CRM.

Entrada: o PPT de modelo usado pelo escritório (ex.: "MODELO DUE - NOVO CRM.pptx").
Saída:   o modelo do CRM, com
  - marcadores {{TOKEN}} no lugar dos valores de exemplo;
  - nomes estáveis nos elementos que o gerador move, esconde ou repete
    (ver docs/DUE-APRESENTACAO-AUTOMATICA.md, seção "Contrato do modelo");
  - um slide genérico de lista de processos (tabela);
  - sem anotações internas (notas e comentários de revisão).

Uso (uv instala o python-pptx sozinho):
  uv run --with python-pptx --with pillow python scripts/due-ppt/prepare-due-template.py \
      "MODELO DUE - NOVO CRM.pptx" "templates/due/MODELO-DUE-CRM.pptx"

Os nomes de origem ("CaixaDeTexto 33", "Tabela 51"...) são do modelo de out/2026.
Se o modelo original mudar, ajuste o mapeamento abaixo; o gerador só depende do
resultado (tokens e nomes).
"""

from __future__ import annotations

import copy
import io
import random
import sys

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

sys.stdout.reconfigure(encoding="utf-8")

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A16_NS = "http://schemas.microsoft.com/office/drawing/2014/main"


# ─── helpers ──────────────────────────────────────────────────────────────────


def iter_shapes(shapes):
    for s in shapes:
        yield s
        if s.shape_type == 6:  # group
            yield from iter_shapes(s.shapes)


def find(slide, name, sid=None):
    found = [
        s
        for s in iter_shapes(slide.shapes)
        if s.name == name and (sid is None or s.shape_id == sid)
    ]
    if len(found) != 1:
        raise SystemExit(f"Esperava 1 shape '{name}' (id={sid}), achei {len(found)}")
    return found[0]


def rename(slide, name, new, sid=None):
    s = find(slide, name, sid)
    s.name = new
    return s


def _clear_runs(p):
    for child in list(p):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p.remove(child)


def _make_run(rPr, text):
    r = etree.SubElement(etree.Element(qn("a:p")), qn("a:r"))
    if rPr is not None:
        r.append(copy.deepcopy(rPr))
    t = etree.SubElement(r, qn("a:t"))
    t.text = text
    return r


def _append_run(p, run):
    end = p.find(qn("a:endParaRPr"))
    if end is not None:
        end.addprevious(run)
    else:
        p.append(run)


def set_text_body(txBody, text):
    """Troca o texto mantendo a formatação do 1º run do 1º parágrafo."""
    ps = txBody.findall(qn("a:p"))
    p0 = ps[0]
    r0 = None
    for p in ps:
        r0 = p.find(qn("a:r"))
        if r0 is not None:
            break
    rPr = r0.find(qn("a:rPr")) if r0 is not None else None
    rPr = copy.deepcopy(rPr) if rPr is not None else None
    for p in ps[1:]:
        txBody.remove(p)
    _clear_runs(p0)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = p0 if i == 0 else copy.deepcopy(p0)
        if i > 0:
            _clear_runs(p)
            txBody.append(p)
        _append_run(p, _make_run(rPr, line))


def set_text(shape, text):
    set_text_body(shape.text_frame._txBody, text)


def tok(slide, name, token, sid=None, new_name=None):
    s = find(slide, name, sid)
    set_text(s, token)
    if new_name:
        s.name = new_name
    return s


def set_runs(shape, texts):
    """Mantém a formatação de cada run (ex.: rótulo em negrito + valor)."""
    p = shape.text_frame._txBody.find(qn("a:p"))
    runs = p.findall(qn("a:r"))
    for i, text in enumerate(texts):
        if i < len(runs):
            runs[i].find(qn("a:t")).text = text
        else:
            extra = copy.deepcopy(runs[-1])
            extra.find(qn("a:t")).text = text
            runs[-1].addnext(extra)
            runs.append(extra)
    for extra in runs[len(texts):]:
        p.remove(extra)


def replace_in_runs(shape, old, new):
    for t in shape.text_frame._txBody.iter(qn("a:t")):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new)
            return
    raise SystemExit(f"'{old}' não encontrado em {shape.name}")


def set_cell(cell, text):
    set_text_body(cell._tc.txBody, text)


def new_a16_ids(el):
    for node in el.iter(f"{{{A16_NS}}}rowId", f"{{{A16_NS}}}colId", f"{{{A16_NS}}}creationId"):
        if node.get("val") is not None:
            if node.tag.endswith("creationId"):
                node.set("id", "{%08X-0000-0000-0000-%012X}" % (random.getrandbits(32), random.getrandbits(48)))
            else:
                node.set("val", str(random.randint(100000000, 4200000000)))


def table_row_template(shape, template_cells, keep_after=None):
    """Deixa: cabeçalho, 1 linha modelo ({{row.*}}) e, se `keep_after`, as N últimas linhas."""
    tbl = shape.table._tbl
    trs = tbl.tr_lst
    keep_tail = trs[len(trs) - keep_after:] if keep_after else []
    tmpl = trs[1]
    for tr in trs[2:]:
        if tr not in keep_tail:
            tbl.remove(tr)
    for i, text in enumerate(template_cells):
        set_text_body(tmpl.tc_lst[i].txBody, text)
    return keep_tail


def copy_shape_to(src_slide, el, dst_slide, rid_cache):
    """Copia um elemento de shape para outro slide, religando imagens/links."""
    new_el = copy.deepcopy(el)
    for node in new_el.iter():
        for attr, val in list(node.attrib.items()):
            if attr.startswith(f"{{{R_NS}}}"):
                if val not in rid_cache:
                    rel = src_slide.part.rels[val]
                    if rel.is_external:
                        rid_cache[val] = dst_slide.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
                    else:
                        rid_cache[val] = dst_slide.part.relate_to(rel.target_part, rel.reltype)
                node.set(attr, rid_cache[val])
    dst_slide.shapes._spTree.append(new_el)
    return new_el


def next_shape_id(slide):
    ids = [int(x) for x in slide.shapes._spTree.xpath(".//p:cNvPr/@id")]
    return max(ids + [0]) + 1


def set_shape_id_name(el, sid, name):
    cnv = el.find(".//" + qn("p:cNvPr"))
    cnv.set("id", str(sid))
    cnv.set("name", name)


def set_creation_id(el, sid):
    """creationId próprio para um shape copiado (o PowerPoint usa no Morph e na coautoria)."""
    for node in el.iter(f"{{{A16_NS}}}creationId"):
        node.set("id", "{%08X-0000-4000-8000-%012X}" % (sid, sid))


def move_slide(prs, slide, new_index):
    lst = prs.slides._sldIdLst
    for sldId in list(lst):
        if prs.slides.part.related_part(sldId.rId) is slide.part:
            lst.remove(sldId)
            lst.insert(new_index, sldId)
            return
    raise SystemExit("slide não encontrado")


def duplicate_slide(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_cache = {}
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        copy_shape_to(src, el, new, rid_cache)
    src_bg = src._element.cSld.bg
    if src_bg is not None:
        new_bg = copy.deepcopy(src_bg)
        new._element.cSld.insert(0, new_bg)
    return new


def strip_notes_and_comments(prs):
    for slide in prs.slides:
        for rId, rel in list(slide.part.rels.items()):
            if rel.reltype == RT.NOTES_SLIDE or "comments" in rel.reltype:
                slide.part.drop_rel(rId)


def placeholder_logo_png():
    img = Image.new("RGBA", (720, 300), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.rounded_rectangle([6, 6, 714, 294], radius=36, outline=(201, 166, 107, 255), width=6)
    try:
        font = ImageFont.truetype("arialbd.ttf", 56)
    except OSError:
        font = ImageFont.load_default()
    text = "LOGO DO CLIENTE"
    w = d.textlength(text, font=font)
    d.text(((720 - w) / 2, 118), text, fill=(201, 166, 107, 255), font=font)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


# ─── preparo ──────────────────────────────────────────────────────────────────


def main(src_path, out_path):
    random.seed(20261020)
    prs = Presentation(src_path)
    S = {i: s for i, s in enumerate(prs.slides, 1)}
    assert len(S) == 24, f"modelo esperado com 24 slides, veio {len(S)}"

    # Marcadores de slide (o gerador acha cada slide pelo nome de um shape SLIDE_*).
    marks = {
        1: ("Título 1", "SLIDE_CAPA"),
        2: ("CaixaDeTexto 17", "SLIDE_SOBRE"),
        3: ("CaixaDeTexto 8", "SLIDE_RECONHECIMENTOS"),
        4: ("CaixaDeTexto 4", "SLIDE_METODOLOGIA"),
        5: ("CaixaDeTexto 3", "SLIDE_DUE_CAPA"),
        6: ("TXT_TITULO", "SLIDE_AREAS"),
        7: ("Imagem 1", "SLIDE_REEST_DIVISOR"),
        8: ("CaixaDeTexto 27", "SLIDE_REEST_RJ"),
        9: ("CaixaDeTexto 27", "SLIDE_REEST_ENDIVIDAMENTO"),
        10: ("CaixaDeTexto 27", "SLIDE_REEST_PAGAMENTO"),
        11: ("CaixaDeTexto 27", "SLIDE_REEST_FORMAS"),
        12: ("CaixaDeTexto 27", "SLIDE_REEST_MEDIDAS"),
        13: ("Imagem 2", "SLIDE_RC_DIVISOR"),
        14: ("CaixaDeTexto 4", "SLIDE_RC_PAINEL"),
        15: ("CaixaDeTexto 6", "SLIDE_DETALHE"),
        16: ("Imagem 8", "SLIDE_CIVEL_DIVISOR"),
        17: ("CaixaDeTexto 6", "SLIDE_CIVEL_PAINEL"),
        18: ("Imagem 14", "SLIDE_TRAB_DIVISOR"),
        19: ("CaixaDeTexto 6", "SLIDE_TRAB_PAINEL"),
        20: ("Imagem 7", "SLIDE_TRIB_DIVISOR"),
        21: ("CaixaDeTexto 6", "SLIDE_TRIB_PAINEL"),
        22: ("TextBox 2", "SLIDE_CONCLUSAO"),
        23: ("TextBox 2", "SLIDE_CONCLUSAO_CREDITO"),
        24: ("CaixaDeTexto 10", "SLIDE_SOCIOS"),
    }

    # ── 4. Metodologia: esferas conforme as áreas da DUE
    replace_in_runs(
        find(S[4], "CaixaDeTexto 11"),
        "abrangendo as esferas cível, tributária, trabalhista e de insolvência.",
        "abrangendo {{METODOLOGIA_ESFERAS}}.",
    )
    find(S[4], "CaixaDeTexto 11").name = "TXT_METODOLOGIA"

    # ── 5. Capa da due: logo do cliente vira placeholder neutro
    s5 = S[5]
    rename(s5, "Elipse 28", "CIRCULO_ESCRITORIO")
    rename(s5, "Elipse 31", "CIRCULO_CLIENTE")
    logo = find(s5, "IMG_LOGO_CLIENTE")
    old_rid = logo._element.blipFill.blip.rEmbed
    _, new_rid = s5.part.get_or_add_image_part(placeholder_logo_png())
    logo._element.blipFill.blip.set(qn("r:embed"), new_rid)
    s5.part.drop_rel(old_rid)

    # ── 6. Áreas envolvidas: todas as áreas como grupos
    s6 = S[6]
    for area in ("REESTRUTURACAO", "TRIBUTARIO"):
        parts = [find(s6, f"AREA_{area}_{suf}") for suf in ("FUNDO", "CARD", "ICONE", "TEXTO")]
        grp = s6.shapes.add_group_shape(parts)
        grp.name = f"AREA_{area}"
    for area in ("CIVEL", "TRABALHISTA", "REC_CRED"):
        find(s6, f"AREA_{area}")

    # ── 8. Reestruturação | RJ
    s8 = S[8]
    for name, token in [
        ("Autos Number", "{{RJ_AUTOS}}"),
        ("Requerentes Value", "{{RJ_RECUPERANDAS}}"),
        ("Competencia Value", "{{RJ_COMPETENCIA}}"),
        ("Juiz Value", "{{RJ_JUIZ}}"),
        ("Relator Value", "{{RJ_RELATOR}}"),
        ("Admin Value", "{{RJ_ADMINISTRADOR}}"),
        ("Summary Text", "{{RJ_RESUMO}}"),
        ("Passivo Value", "{{RJ_PASSIVO_TOTAL}}"),
        ("Status Chip Text", "{{RJ_DEFERIMENTO}}"),
        ("Ajuizamento Value", "{{RJ_AJUIZAMENTO}}"),
        ("Foro Value", "{{RJ_FORO}}"),
    ]:
        tok(s8, name, token)

    # ── 9. Composição do endividamento
    s9 = S[9]
    t9 = rename(s9, "Tabela 102", "TBL_RJ_CLASSES")
    tail = table_row_template(
        t9, ["{{row.classe}}", "{{row.edital}}", "{{row.aj}}", "{{row.variacao}}"], keep_after=2
    )
    for tc, text in zip(tail[0].tc_lst[1:], ["{{RJ_TOTAL_EDITAL}}", "{{RJ_TOTAL_AJ}}", "{{RJ_TOTAL_VARIACAO}}"]):
        set_text_body(tc.txBody, text)
    for tc, text in zip(tail[1].tc_lst[1:], ["{{RJ_EXTRA_EDITAL}}", "{{RJ_EXTRA_AJ}}", "{{RJ_EXTRA_VARIACAO}}"]):
        set_text_body(tc.txBody, text)
    for name, token in [
        ("Text 105", "{{RJ_CONCURSAL_AJ}}"),
        ("Text 109", "{{RJ_VARIACAO_EDITAL}}"),
        ("Text 111", "{{RJ_EXTRACONCURSAL_AJ}}"),
        ("Text 113", "{{RJ_ENDIVIDAMENTO_TOTAL}}"),
        ("Text 116", "{{RJ_ENDIVIDAMENTO_RESUMO}}"),
    ]:
        tok(s9, name, token)

    # ── 10. Forma de pagamento
    t10 = rename(S[10], "Tabela 102", "TBL_RJ_PAGAMENTO")
    table_row_template(
        t10,
        ["{{row.classe}}", "{{row.desagio}}", "{{row.carencia}}", "{{row.forma}}", "{{row.atualizacao}}"],
    )

    # ── 11. Formas especiais
    for name, token in [
        ("Text 102", "{{RJ_OPCAO_A_ROTULO}}"),
        ("Text 103", "{{RJ_OPCAO_A_TITULO}}"),
        ("Text 106", "{{RJ_OPCAO_A_CONTRAPARTIDA}}"),
        ("Text 109", "{{RJ_OPCAO_A_CONDICOES}}"),
        ("Text 110", "{{RJ_OPCAO_A_OBS}}"),
        ("Text 112", "{{RJ_OPCAO_B_ROTULO}}"),
        ("Text 113", "{{RJ_OPCAO_B_TITULO}}"),
        ("Text 116", "{{RJ_OPCAO_B_CONTRAPARTIDA}}"),
        ("Text 119", "{{RJ_OPCAO_B_CONDICOES}}"),
        ("Text 120", "{{RJ_OPCAO_B_OBS}}"),
    ]:
        tok(S[11], name, token)

    # ── 12. Principais medidas
    for n, base in ((1, 102), (2, 112), (3, 122)):
        for off, field in ((0, "CATEGORIA"), (1, "TITULO"), (4, "DESTAQUE"), (6, "PREVE"), (8, "OBJETIVO")):
            tok(S[12], f"Text {base + off}", f"{{{{RJ_MEDIDA_{n}_{field}}}}}")

    # ── 14. Recuperação de crédito — painel
    s14 = S[14]
    for name, first, token in [
        ("CaixaDeTexto 15", " TOTAL DE PROCESSOS", "{{RC_TOTAL}}"),
        ("CaixaDeTexto 25", "VALOR DE CAUSA", "{{RC_VALOR_CAUSA}}"),
        ("CaixaDeTexto 29", "MÉDIA DE VALOR DE CAUSA", "{{RC_MEDIA}}"),
    ]:
        shape = find(s14, name)
        ps = shape.text_frame._txBody.findall(qn("a:p"))
        # 2º parágrafo = valor; mantém o 1º (rótulo) intacto
        p1 = ps[1]
        runs = p1.findall(qn("a:r"))
        runs[0].find(qn("a:t")).text = token
        for extra in runs[1:]:
            p1.remove(extra)
    tok(s14, "CaixaDeTexto 3", "{{RC_ATIVOS}}")
    tok(s14, "CaixaDeTexto 8", "{{RC_ARQUIVADOS}}")
    rename(s14, "Gráfico 7", "CHART_RC_TIPO")
    for i, name in enumerate(
        [
            "Retângulo: Cantos Arredondados 12",
            "Retângulo: Cantos Arredondados 14",
            "Retângulo: Cantos Arredondados 18",
            "Retângulo: Cantos Arredondados 21",
            "Retângulo: Cantos Arredondados 24",
        ],
        1,
    ):
        tok(s14, name, f"{{{{RC_TIPO_QTD_{i}}}}}", new_name=f"RC_TIPO_QTD_{i}")
    rename(s14, "Gráfico 47", "CHART_RC_FAIXA")
    tok(s14, "CaixaDeTexto 48", "{{RC_FAIXA_TOTAL}}")
    for i, name in enumerate(["CaixaDeTexto 52", "CaixaDeTexto 54", "CaixaDeTexto 56", "CaixaDeTexto 58", "CaixaDeTexto 60"], 1):
        tok(s14, name, f"{{{{RC_FAIXA_QTD_{i}}}}}")
    rename(s14, "Gráfico 63", "CHART_RC_ANO")
    rename(s14, "Gráfico 73", "CHART_RC_TEMPO")
    for i, name in enumerate(["CaixaDeTexto 75", "CaixaDeTexto 77", "CaixaDeTexto 79", "CaixaDeTexto 81"], 1):
        tok(s14, name, f"{{{{RC_TEMPO_QTD_{i}}}}}")
    tok(s14, "CaixaDeTexto 96", "{{RC_TEMPO_MEDIO}}")
    t14 = rename(s14, "Tabela 84", "TBL_RC_TOP5")
    table_row_template(t14, ["{{row.pos}}", "{{row.devedor}}", "{{row.processo}}", "{{row.valor}}"])
    tok(s14, "CaixaDeTexto 86", "{{RC_TOP5_TITULO}}")
    tok(s14, "CaixaDeTexto 88", "Total dos {{RC_TOP5_N}}")
    tok(s14, "CaixaDeTexto 90", "{{RC_TOP5_TOTAL}}")
    tok(s14, "CaixaDeTexto 92", "({{RC_TOP5_PCT}} do valor total da carteira)")

    # ── 15. Detalhe de processo (genérico: um slide por processo, qualquer área)
    s15 = S[15]
    tok(s15, "CaixaDeTexto 6", "{{DET_AREA}}")
    for name, token in [
        ("Processo Number", "{{DET_PROCESSO}}"),
        ("Reclamante Label", "{{DET_L_PARTE_A}}"),
        ("Reclamante Value", "{{DET_PARTE_A}}"),
        ("Reclamadas Label", "{{DET_L_PARTE_B}}"),
        ("Reclamadas Value", "{{DET_PARTE_B}}"),
        ("Pedidos Label", "{{DET_L_BLOCO_1}}"),
        ("Pedidos Value", "{{DET_BLOCO_1}}"),
        ("Alegacoes Label", "{{DET_L_BLOCO_2}}"),
        ("Alegacoes Value", "{{DET_BLOCO_2}}"),
        ("Passivo Label", "{{DET_L_VALOR}}"),
        ("Passivo Value", "{{DET_VALOR}}"),
        ("Status Chip", "{{DET_STATUS}}"),
        ("Vara Value", "{{DET_VARA}}"),
        ("Valor Value", "{{DET_TRIBUNAL}}"),
        ("Distrib Value", "{{DET_AJUIZAMENTO}}"),
        ("Contrato Label", "{{DET_L_INFO_1}}"),
        ("Contrato Value", "{{DET_INFO_1}}"),
        ("Cargo Label", "{{DET_L_INFO_2}}"),
        ("Cargo Value", "{{DET_INFO_2}}"),
        ("Remun Label", "{{DET_L_INFO_3}}"),
        ("Remun Value", "{{DET_INFO_3}}"),
        ("Resumo Value", "{{DET_RESUMO}}"),
    ]:
        tok(s15, name, token)
    # Pedidos/Alegações viram 1 parágrafo; o gerador expande linhas com **negrito**.

    # ── 17. Cível — painel
    s17 = S[17]
    for name, token in [
        ("CaixaDeTexto 33", "{{CIVEL_PASSIVO}}"),
        ("CaixaDeTexto 35", "{{CIVEL_ATIVOS}}"),
        ("CaixaDeTexto 37", "{{CIVEL_SUSPENSOS}}"),
        ("CaixaDeTexto 52", "{{CIVEL_MEDIA}}"),
        ("CaixaDeTexto 54", "{{CIVEL_MAIOR}}"),
        ("CaixaDeTexto 177", "{{CIVEL_LEITURA}}"),
        ("CaixaDeTexto 63", "{{CIVEL_TOP5_TOTAL}}"),
    ]:
        tok(s17, name, token)
    tok(s17, "CaixaDeTexto 59", "Total dos {{CIVEL_TOP5_N}}")
    tok(s17, "CaixaDeTexto 55", "{{CIVEL_TOP5_TITULO}}")
    slots17 = {
        1: ("Retângulo: Cantos Arredondados 106", "Elipse 105", "CaixaDeTexto 108", "CaixaDeTexto 110", "CaixaDeTexto 112", "Conector reto 146"),
        2: ("Retângulo: Cantos Arredondados 114", "Elipse 126", "CaixaDeTexto 118", "CaixaDeTexto 120", "CaixaDeTexto 122", "Conector reto 149"),
    }
    name_fase_slots(s17, "CIVEL_FASE", slots17)
    rename(s17, "Conector reto 151", "CIVEL_FASE_LINHA")
    # sócio no polo passivo: slot 1 + cópia como slot 2
    socio1 = ("Retângulo: Cantos Arredondados 9", "Elipse 11", "CaixaDeTexto 13", "CaixaDeTexto 15", "CaixaDeTexto 17", "Conector reto 29")
    name_fase_slots(s17, "CIVEL_SOCIO", {1: socio1})
    rename(s17, "Conector reto 1", "CIVEL_SOCIO_LINHA")
    dy = Emu(Inches(1.42))
    sid = next_shape_id(s17)
    for suf in ("BARRA", "PONTO", "NOME", "QTD", "VALOR", "CONECTOR"):
        src = find(s17, f"CIVEL_SOCIO_1_{suf}")
        el = copy.deepcopy(src._element)
        s17.shapes._spTree.append(el)
        set_shape_id_name(el, sid, f"CIVEL_SOCIO_2_{suf}")
        set_creation_id(el, sid)
        sid += 1
        off = el.find(".//" + qn("a:off"))
        off.set("y", str(int(off.get("y")) + dy))
        for t in el.iter(qn("a:t")):
            t.text = (t.text or "").replace("CIVEL_SOCIO_1", "CIVEL_SOCIO_2")
    t17 = rename(s17, "Tabela 51", "TBL_CIVEL_TOP5")
    table_row_template(t17, ["{{row.parte}}", "{{row.processo}}", "{{row.valor}}"])
    set_cell(t17.table.cell(0, 0), "{{CIVEL_TOP5_L_PARTE}}")
    # Processos por região: até 4 colunas (sigla + quantidade) com divisórias. O gerador
    # mantém as que usa e divide a faixa entre elas.
    for i, (nome, qtd) in {1: ("CaixaDeTexto 66", "CaixaDeTexto 56"), 2: ("CaixaDeTexto 68", "CaixaDeTexto 58")}.items():
        tok(s17, nome, f"{{{{CIVEL_TRIBUNAL_{i}}}}}", new_name=f"CIVEL_TRIBUNAL_{i}_NOME")
        tok(s17, qtd, f"{{{{CIVEL_TRIBUNAL_{i}_QTD}}}}", new_name=f"CIVEL_TRIBUNAL_{i}_QTD")
    rename(s17, "Conector reto 194", "CIVEL_TRIBUNAL_DIV_1")
    sid = next_shape_id(s17)
    copias = [(f"CIVEL_TRIBUNAL_2_{suf}", f"CIVEL_TRIBUNAL_{i}_{suf}") for i in (3, 4) for suf in ("NOME", "QTD")]
    copias += [("CIVEL_TRIBUNAL_DIV_1", f"CIVEL_TRIBUNAL_DIV_{i}") for i in (2, 3)]
    for src, dst in copias:
        el = copy.deepcopy(find(s17, src)._element)
        s17.shapes._spTree.append(el)
        set_shape_id_name(el, sid, dst)
        set_creation_id(el, sid)
        sid += 1
        for t in el.iter(qn("a:t")):
            t.text = (t.text or "").replace("CIVEL_TRIBUNAL_2", dst.rsplit("_", 1)[0])
    col_x, col_w = Inches(10.12), Inches(2.92) // 4
    for i in range(1, 5):
        for suf in ("NOME", "QTD"):
            shape = find(s17, f"CIVEL_TRIBUNAL_{i}_{suf}")
            shape.left, shape.width = Emu(col_x + (i - 1) * col_w), Emu(col_w)
            shape.text_frame.margin_left = shape.text_frame.margin_right = 0
    for i in range(1, 4):
        div = find(s17, f"CIVEL_TRIBUNAL_DIV_{i}")
        div.left = Emu(col_x + i * col_w)

    # ── 19. Trabalhista — painel
    s19 = S[19]
    for name, token in [
        ("CaixaDeTexto 33", "{{TRAB_PASSIVO}}"),
        ("CaixaDeTexto 35", "{{TRAB_ATIVOS}}"),
        ("CaixaDeTexto 37", "{{TRAB_ARQUIVADOS}}"),
        ("CaixaDeTexto 52", "{{TRAB_MEDIA}}"),
        ("CaixaDeTexto 54", "{{TRAB_MAIOR}}"),
        ("CaixaDeTexto 160", "{{TRAB_OBS_ANO}}"),
        ("CaixaDeTexto 171", "{{TRAB_OBS_PEDIDOS}}"),
        ("CaixaDeTexto 177", "{{TRAB_LEITURA}}"),
    ]:
        tok(s19, name, token)
    slots19 = {
        1: ("Retângulo: Cantos Arredondados 106", "Elipse 105", "CaixaDeTexto 108", "CaixaDeTexto 110", "CaixaDeTexto 112", "Conector reto 146"),
        2: ("Retângulo: Cantos Arredondados 114", "Elipse 116", "CaixaDeTexto 118", "CaixaDeTexto 120", "CaixaDeTexto 122", "Conector reto 148"),
        3: ("Retângulo: Cantos Arredondados 124", "Elipse 126", "CaixaDeTexto 128", "CaixaDeTexto 130", "CaixaDeTexto 132", "Conector reto 149"),
    }
    name_fase_slots(s19, "TRAB_FASE", slots19)
    rename(s19, "Conector reto 151", "TRAB_FASE_LINHA")
    rename(s19, "Chart", "CHART_TRAB_PEDIDOS", sid=304)
    rename(s19, "Chart", "CHART_TRAB_ANO", sid=305)

    # ── 21. Tributário — painel
    s21 = S[21]
    for name, token in [
        ("CaixaDeTexto 33", "{{TRIB_PROCESSOS}}"),
        ("CaixaDeTexto 50", "{{TRIB_PASSIVO_JUDICIAL}}"),
        ("CaixaDeTexto 61", "{{TRIB_TRAMITACAO}}"),
        ("CaixaDeTexto 65", "{{TRIB_SUSPENSOS}}"),
        ("CaixaDeTexto 69", "{{TRIB_ADMINISTRATIVO}}"),
        ("CaixaDeTexto 177", "{{TRIB_PONTOS}}"),
        ("CaixaDeTexto 116", "{{TRIB_LEITURA}}"),
    ]:
        tok(s21, name, token)
    t_origem = rename(s21, "Tabela 73", "TBL_TRIB_ORIGEM")
    tail = table_row_template(t_origem, ["{{row.origem}}", "{{row.valor}}"], keep_after=1)
    set_text_body(tail[0].tc_lst[1].txBody, "{{TRIB_ORIGEM_TOTAL}}")
    t_nat = rename(s21, "Tabela 84", "TBL_TRIB_NATUREZA")
    for row, token in zip(range(1, 5), ["{{TRIB_PGFN_TRIBUTARIO}}", "{{TRIB_PGFN_PREVIDENCIARIO}}", "{{TRIB_PGFN_MULTA}}", "{{TRIB_PGFN_TOTAL}}"]):
        set_cell(t_nat.table.cell(row, 1), token)
    t_adm = rename(s21, "Tabela 99", "TBL_TRIB_ADM")
    for row, token in zip(range(1, 4), ["{{TRIB_ADM_PGFN}}", "{{TRIB_ADM_PGE}}", "{{TRIB_ADM_TOTAL}}"]):
        set_cell(t_adm.table.cell(row, 1), token)

    # ── 22. Conclusão (sem crédito)
    s22 = S[22]
    rows22 = {
        "REEST": ("Reest 600", "Reest 601", "Reest 602", "Reest 603", "Reest 604", "Gráfico 9"),
        "TRIB": ("Trib 16", "Trib 17", "Trib 204", "Trib 205", "Trib 206", "Gráfico 16"),
        "TRAB": ("TextBox 15", "TextBox 16", "Shape 204", "Shape 205", "TextBox 206", "Gráfico 52"),
        "CIVEL": ("Civel 16", "Civel 17", "Civel 204", "Civel 205", "Civel 206", "Gráfico 38"),
    }
    name_conclusao_rows(s22, rows22)
    rename(s22, "Rectangle 14", "CONC_SEPARADOR")
    tok(s22, "TextBox 19", "{{CONC_TOTAL}}")
    tok(s22, "TextBox 208", "{{CONC_DESCRICAO}}")
    tok(s22, "Footnote", "{{CONC_NOTAS}}", new_name="CONC_NOTAS")

    # ── 23. Conclusão (com recuperação de crédito): traz Reestruturação/Tributário do 22
    s23 = S[23]
    rows23 = {
        "TRAB": ("TextBox 15", "TextBox 16", "Shape 204", "Shape 205", "TextBox 206", "Gráfico 35"),
        "CIVEL": ("Civel 16", "Civel 17", "Civel 204", "Civel 205", "Civel 206", "Graphic 305"),
    }
    name_conclusao_rows(s23, rows23)
    rid_cache = {}
    sid = next_shape_id(s23)
    trab_valor = find(s23, "CONC_TRAB_VALOR")
    trab_trilho = find(s23, "CONC_TRAB_TRILHO")
    trab_pct = find(s23, "CONC_TRAB_PCT")
    for area in ("REEST", "TRIB"):
        for suf in ("NOME", "VALOR", "TRILHO", "BARRA", "PCT", "ICONE"):
            src = find(s22, f"CONC_{area}_{suf}")
            el = copy_shape_to(s22, src._element, s23, rid_cache)
            set_shape_id_name(el, sid, f"CONC_{area}_{suf}")
            sid += 1
            shape = find(s23, f"CONC_{area}_{suf}")
            if suf == "VALOR":
                shape.left, shape.width = trab_valor.left, trab_valor.width
            elif suf in ("TRILHO", "BARRA"):
                shape.width = trab_trilho.width if suf == "TRILHO" else Emu(int(trab_trilho.width * 0.5))
            elif suf == "PCT":
                shape.width = trab_pct.width
                set_text(shape, f"{{{{CONC_{area}_PCT}}}}")
    # Linhas copiadas começam na mesma altura da 1ª linha (o gerador empilha).
    def row_top(prefix):
        return min(s.top for s in iter_shapes(s23.shapes) if s.name.startswith(prefix))
    trab_top = row_top("CONC_TRAB_")
    for area in ("REEST", "TRIB"):
        dy = trab_top - row_top(f"CONC_{area}_")
        for s in iter_shapes(s23.shapes):
            if s.name.startswith(f"CONC_{area}_"):
                s.top = Emu(s.top + dy)
    rename(s23, "Rectangle 14", "CONC_SEPARADOR")
    tok(s23, "TextBox 19", "{{CONC_TOTAL}}")
    tok(s23, "TextBox 208", "{{CONC_DESCRICAO}}")
    tok(s23, "Credit Total Value", "{{CRED_TOTAL}}")
    tok(s23, "Credit Label 1", "{{CRED_DESCRICAO}}")

    # ── Slide de lista de processos (tabela), criado a partir do painel Cível
    lista = duplicate_slide(prs, s17)
    keep = {
        "Retângulo 2",
        "Retângulo 8",
        "Retângulo: Cantos Arredondados 88",
        "Elipse 90",
        "Gráfico 181",
        "CaixaDeTexto 104",
        "Retângulo: Cantos Arredondados 10",
        "Elipse 19",
        "Gráfico 193",
        "CaixaDeTexto 26",
        "CaixaDeTexto 33",
        "Conector reto 49",
        "TBL_CIVEL_TOP5",
        "CaixaDeTexto 6",
    }
    for shp in list(lista.shapes):
        if shp.name not in keep:
            shp._element.getparent().remove(shp._element)

    def place(name, x, y, w=None, h=None, new_name=None, text=None):
        shp = find(lista, name)
        shp.left, shp.top = Emu(Inches(x)), Emu(Inches(y))
        if w is not None:
            shp.width = Emu(Inches(w))
        if h is not None:
            shp.height = Emu(Inches(h))
        if text is not None:
            set_text(shp, text)
        if new_name:
            shp.name = new_name
        return shp

    tok(lista, "CaixaDeTexto 6", "{{LISTA_AREA}}", new_name="SLIDE_LISTA")
    place("Retângulo: Cantos Arredondados 88", 0.17, 0.98, 12.99, 6.36)
    place("Elipse 90", 0.30, 1.10, 0.56, 0.56)
    place("Gráfico 181", 0.36, 1.16, 0.44, 0.44)
    place("CaixaDeTexto 104", 0.95, 1.20, 6.0, 0.32, text="{{LISTA_TITULO}}")
    place("Retângulo: Cantos Arredondados 10", 8.10, 0.12, 5.06, 0.70)
    place("Elipse 19", 8.20, 0.19, 0.56, 0.56)
    place("Gráfico 193", 8.27, 0.26, 0.42, 0.42)
    place("CaixaDeTexto 26", 8.82, 0.31, 2.10, 0.30, text="{{LISTA_L_VALOR}}")
    place("Conector reto 49", 10.86, 0.22, h=0.50)
    place("CaixaDeTexto 33", 10.92, 0.25, 2.20, 0.40, text="{{LISTA_VALOR}}")
    tbl_shape = place("TBL_CIVEL_TOP5", 0.40, 1.78, 12.53, None, new_name="TBL_LISTA")
    tbl = tbl_shape.table._tbl
    grid = tbl.tblGrid
    gridcols = grid.findall(qn("a:gridCol"))
    while len(gridcols) < 5:
        gc = copy.deepcopy(gridcols[-1])
        grid.append(gc)
        gridcols.append(gc)
    for gc, w in zip(gridcols, (2.15, 2.85, 2.85, 1.55, 3.13)):
        gc.set("w", str(Emu(Inches(w))))
    for tr in tbl.tr_lst:
        tcs = tr.tc_lst
        while len(tcs) < 5:
            tc = copy.deepcopy(tcs[-1])
            tcs[-1].addnext(tc)
            tcs = tr.tc_lst
    new_a16_ids(tbl)
    header = ["Processo", "{{LISTA_L_PARTE_A}}", "{{LISTA_L_PARTE_B}}", "{{LISTA_L_VALOR_COL}}", "{{LISTA_L_INFO}}"]
    for i, text in enumerate(header):
        set_text_body(tbl.tr_lst[0].tc_lst[i].txBody, text)
    for i, text in enumerate(["{{row.processo}}", "{{row.parteA}}", "{{row.parteB}}", "{{row.valor}}", "{{row.info}}"]):
        set_text_body(tbl.tr_lst[1].tc_lst[i].txBody, text)
    move_slide(prs, lista, 15)  # logo depois do slide de detalhe

    # marcadores de slide (depois dos tokens: alguns títulos foram usados acima)
    for idx, (name, mark) in marks.items():
        rename(S[idx], name, mark)

    strip_notes_and_comments(prs)
    prs.save(out_path)
    print("ok ->", out_path)


def name_fase_slots(slide, prefix, slots):
    for i, (barra, ponto, nome, qtd, valor, conector) in slots.items():
        rename(slide, barra, f"{prefix}_{i}_BARRA")
        rename(slide, ponto, f"{prefix}_{i}_PONTO")
        tok(slide, nome, f"{{{{{prefix}_{i}}}}}", new_name=f"{prefix}_{i}_NOME")
        shape = rename(slide, qtd, f"{prefix}_{i}_QTD")
        set_runs(shape, ["Processos: ", f"{{{{{prefix}_{i}_QTD}}}}"])
        shape = rename(slide, valor, f"{prefix}_{i}_VALOR")
        set_runs(shape, ["Valor envolvido: ", f"{{{{{prefix}_{i}_VALOR}}}}"])
        rename(slide, conector, f"{prefix}_{i}_CONECTOR")


def name_conclusao_rows(slide, rows):
    for area, (nome, valor, trilho, barra, pct, icone) in rows.items():
        rename(slide, nome, f"CONC_{area}_NOME")
        tok(slide, valor, f"{{{{CONC_{area}_VALOR}}}}", new_name=f"CONC_{area}_VALOR")
        rename(slide, trilho, f"CONC_{area}_TRILHO")
        rename(slide, barra, f"CONC_{area}_BARRA")
        tok(slide, pct, f"{{{{CONC_{area}_PCT}}}}", new_name=f"CONC_{area}_PCT")
        rename(slide, icone, f"CONC_{area}_ICONE")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
